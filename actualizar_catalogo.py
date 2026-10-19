"""Convierte productos.xlsx en datos.js (lo que lee catalogo.html).
Uso: python actualizar_catalogo.py   (requiere: pip install openpyxl)"""
import json, os, sys
from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(AQUI, "productos.xlsx"), data_only=True)

def filas(hoja):
    ws = wb[hoja]
    cab = [str(c.value).strip() if c.value else "" for c in ws[1]]
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in fila):
            yield {cab[i]: ("" if v is None else v) for i, v in enumerate(fila) if cab[i]}

productos, avisos = [], []
for p in filas("Productos"):
    if str(p.get("publicado", "")).strip().lower() not in ("sí", "si", "yes", "x"):
        continue
    for campo in ("id", "nombre", "categoria", "precio"):
        if p.get(campo) in ("", None):
            avisos.append(f"Fila {p.get('id') or p.get('nombre') or '?'}: falta '{campo}'")
    try:
        p["precio"] = float(p["precio"])
    except (TypeError, ValueError):
        avisos.append(f"{p.get('id')}: el precio no es un número"); p["precio"] = 0
    img = str(p.get("imagen", "")).strip()
    if img and not os.path.exists(os.path.join(AQUI, img)):
        avisos.append(f"{p.get('id')}: no existe el archivo {img}")
    p.pop("publicado", None)
    productos.append({k: (str(v).strip() if isinstance(v, str) else v) for k, v in p.items()})

faq = [{"p": f["pregunta"], "r": f["respuesta"]} for f in filas("FAQ") if f.get("pregunta")]
if "Config" in wb.sheetnames:
    cfg = {str(c["clave"]).strip(): str(c["valor"]).strip() for c in filas("Config") if c.get("clave")}
    moneda, aviso = cfg.get("moneda", "S/"), cfg.get("aviso", "")
else:
    moneda, aviso = "S/", "Prototipo con datos ficticios"
datos = {"moneda": moneda, "aviso": aviso, "productos": productos, "faq": faq}
with open(os.path.join(AQUI, "datos.js"), "w", encoding="utf-8") as f:
    f.write("window.CATALOGO = " + json.dumps(datos, ensure_ascii=False, indent=1) + ";\n")
print(f"OK: {len(productos)} productos publicados, {len(faq)} preguntas frecuentes.")
for a in avisos: print("AVISO:", a)
